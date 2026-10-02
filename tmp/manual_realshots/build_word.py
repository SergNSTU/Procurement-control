from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).parent
OUT=ROOT.parent.parent/'output'/'word'/'Procurement_Control_WPF_2.0_User_Guide_RU.docx'
SHOTS=ROOT/'screens'
OUT.parent.mkdir(parents=True,exist_ok=True)
d=Document(); sec=d.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.7); sec.right_margin=Inches(.7)
styles=d.styles; styles['Normal'].font.name='Arial'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Arial'); styles['Normal'].font.size=Pt(10)
for name,size in [('Title',26),('Heading 1',18),('Heading 2',14)]:
 st=styles[name];st.font.name='Arial';st._element.rPr.rFonts.set(qn('w:eastAsia'),'Arial');st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
def add_page_number(par):
 run=par.add_run();fldChar1=OxmlElement('w:fldChar');fldChar1.set(qn('w:fldCharType'),'begin');instrText=OxmlElement('w:instrText');instrText.set(qn('xml:space'),'preserve');instrText.text=' PAGE ';fldChar2=OxmlElement('w:fldChar');fldChar2.set(qn('w:fldCharType'),'end');run._r.append(fldChar1);run._r.append(instrText);run._r.append(fldChar2)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;footer.add_run('Procurement Control WPF 2.0 | ');add_page_number(footer)
def p(text='',bold=False):
 q=d.add_paragraph();r=q.add_run(text);r.bold=bold;return q
def bullets(items):
 for x in items:d.add_paragraph(x,style='List Bullet')
def fig(name,caption):
 d.add_picture(str(SHOTS/name),width=Inches(6.85));q=d.add_paragraph(caption);q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.runs[0].italic=True;q.runs[0].font.size=Pt(9)
def page():d.add_page_break()

q=d.add_paragraph(style='Title');q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.add_run('Procurement Control WPF 2.0')
q=d.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.add_run('Руководство пользователя').bold=True;q.runs[0].font.size=Pt(18)
d.add_paragraph('Для менеджеров по закупкам',style='Subtitle').alignment=WD_ALIGN_PARAGRAPH.CENTER
d.add_paragraph('\n\n\nТекущая рабочая база\n8 сентября 2026 г.').alignment=WD_ALIGN_PARAGRAPH.CENTER
page()
d.add_heading('Содержание',0)
p('1. Быстрый старт\n2. Сравнение RRFQ\n3. Компэл Парсер\n4. Контроль закупки и связанные разделы\n5. Аналитика и база квот\n6. Настройки, задачи и безопасная работа')
d.add_heading('Как читать руководство',1)
p('Снимки сделаны из актуального WPF-интерфейса на текущей рабочей базе. Они показывают реальные поля, состояния и состав данных на момент подготовки документа.')
page()
d.add_heading('1. Быстрый старт',0)
p('Procurement Control объединяет сравнение предложений поставщиков, контроль закупки, документы, задачи, историю и аналитику. Основная навигация находится слева, активный раздел открывается справа.')
d.add_heading('Перед началом',1)
bullets(['Запускайте WPF 2.0 из папки переносимой сборки. Не запускайте приложение из ZIP-архива.','Для сравнения RRFQ и создания Excel-файлов требуется Microsoft Excel.','Перед массовыми изменениями создавайте резервную копию в разделе «Настройки».','Проверяйте включенные позиции и выбранного победителя перед созданием итогового RRFQ.'])
page()
d.add_heading('2. Сравнение RRFQ',0)
fig('01_rrfq_real.png','Рисунок 1. Текущий экран «Сравнение RRFQ».')
d.add_heading('Подготовка сравнения',1)
bullets(['Нажмите «Выбрать...» напротив RFQ и укажите исходную книгу .xlsx или .xls.','Нажмите «Добавить поставщиков» и выделите ответы поставщиков. Можно добавить несколько файлов одновременно.','Проверьте и при необходимости исправьте название в колонке «Поставщик».','Удалите ошибочно добавленные строки и выберите приоритет: «Цена» или «Срок».','Нажмите «Сформировать preview».'])
d.add_heading('Preview и корректировка',1)
bullets(['Используйте фильтры «Все строки», «Только победители», «Без квот», «Предупреждения» и «Ручная проверка».','Выделите позицию, чтобы увидеть все найденные квоты в нижней таблице.','Для переопределения автоматического выбора нажмите «Сделать победителем».','Для отсутствующей или несопоставленной позиции примените «Ручное сопоставление» либо «Добавить квоту вручную».','После проверки нажмите «Создать результирующий RRFQ».'])
page()
d.add_heading('3. Компэл Парсер',0)
fig('08_CompelParser.png','Рисунок 2. Текущий экран «Компэл Парсер».')
p('Парсер читает сохраненный HTML расчета SDS Компэл, показывает позиции и позволяет создать Excel или добавить Компэл в состав сравнения RRFQ.')
d.add_heading('Последовательность работы',1)
bullets(['Сохраните страницу расчета SDS Компэл как HTML (.html или .htm). Не используйте PDF или скриншот страницы.','Нажмите «Выбрать HTML». Разбор запускается автоматически; «Разобрать» повторяет обработку выбранного файла.','Проверьте карточки: строки всего, строки с ценой, строки без предложения и USD-итоги.','При необходимости выгрузите проверочную книгу кнопкой «Выгрузить Excel».','Нажмите «В сравнение RRFQ», затем добавьте другие предложения и выполните общий preview.'])
d.add_heading('Если нет результата',1)
bullets(['Проверьте расширение и убедитесь, что HTML содержит расчет, а не страницу входа или ошибку сайта.','Повторно сохраните страницу после полной загрузки.','При проблемах с текстом пересохраните HTML в UTF-8 или Windows-1251.','Строку без цены обработайте вручную в RRFQ либо исключите из результата.'])
page()
d.add_heading('4. Контроль закупки',0)
fig('04_Purchase.png','Рисунок 3. «Контроль закупки» на текущей базе.')
p('Верхние карточки показывают количество активных и просроченных сделок, задачи в работе и объекты, требующие внимания. Основная таблица содержит текущие сделки и их операционные статусы.')
d.add_heading('Ежедневная работа',1)
bullets(['Нажмите «Новая сделка» для создания записи.','Выберите сделку в таблице, чтобы заполнить этап, приоритет, напоминание и комментарий.','Добавляйте поставщиков, суммы PI, оплату, сроки и документы в нижней части раздела.','Используйте «Свести RFQ» для подготовки сравнения из документов выбранной сделки.','Проверяйте индикаторы «Оплачено», «Готово» и «Отслеживание» перед закрытием сделки.'])
page()
d.add_heading('5. Аналитика и база квот',0)
fig('02_SupplierAnalytics.png','Рисунок 4. «Аналитика поставщиков» на текущей базе.')
p('Аналитика использует сохраненные результаты сравнений. Фильтруйте по производителю, поставщику и датам, чтобы получать рекомендации и статистику по выборкам.')
fig('07_QuoteBase.png','Рисунок 5. «База квот» на текущей базе.')
bullets(['База квот хранит все найденные предложения, а не только победителей.','Перед повторным запросом используйте поиск и фильтры по PN, поставщику и сроку.','Проверяйте дату, валюту, MOQ и статус перед использованием исторической цены.'])
page()
d.add_heading('6. Задачи и настройки',0)
fig('05_Components.png','Рисунок 6. «Задачи» на текущей базе.')
p('В задачах фиксируйте сделку, этап, описание, следующий шаг, дедлайн и приоритет. После завершения актуализируйте статус и напоминание.')
fig('11_Settings.png','Рисунок 7. «Настройки» и резервные копии.')
d.add_heading('Сохранность данных',1)
bullets(['Проверяйте папку документов и сохраняйте настройки после изменения пути.','Перед изменением большого числа сделок создавайте резервную копию.','Используйте восстановление только после проверки состава копии.','Не удаляйте элементы безвозвратно, пока не проверили, что восстановление не понадобится.'])
d.add_heading('Типовые ошибки',1)
bullets(['Не выбран RFQ: укажите одну исходную книгу.','Нет файлов поставщиков: добавьте как минимум один файл.','Preview пустой: проверьте заголовки и наличие строк в Excel.','Нет доступа к файлу: закройте его в Excel и проверьте право записи.'])
d.save(OUT)
print(OUT)
