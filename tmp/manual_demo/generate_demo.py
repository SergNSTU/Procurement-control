from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).parent / "input"
ROOT.mkdir(parents=True, exist_ok=True)

HEADERS = [
    "Value", "Russian remark", "China remark", "PN", "D/C", "Mfg from Russia",
    "Mfg from China", "Q-ty in packing", "Q-ty to Buy/pcs", "Unit price",
    "Currency", "Lead time", "Supplier",
]

RFQ_ROWS = [
    ["MCU-10A demo controller", "Требуется аналог", "", "MCU-10A", "", "", "", 1, 100, "", "", "", ""],
    ["SENSOR-5V demo sensor", "", "", "SENSOR-5V", "", "", "", 1, 50, "", "", "", ""],
    ["CAP-100U demo capacitor", "", "", "CAP-100U", "", "", "", 10, 200, "", "", "", ""],
]

def make_book(name, rows):
    book = Workbook()
    ws = book.active
    ws.title = "RRFQ"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in rows:
        ws.append(row)
    for column in ws.columns:
        letter = column[0].column_letter
        ws.column_dimensions[letter].width = min(30, max(12, max(len(str(c.value or "")) for c in column) + 2))
    path = ROOT / name
    book.save(path)
    return path

make_book("RFQ_DEMO.xlsx", RFQ_ROWS)
make_book("RRFQ_Atlas_Demo.xlsx", [
    ["MCU-10A demo controller", "", "", "MCU-10A", "", "", "Nordic", 1, 100, 11.20, "USD", "14 дней", "Атлас Электроника"],
    ["SENSOR-5V demo sensor", "", "", "SENSOR-5V", "", "", "Sensoria", 1, 50, 4.80, "USD", "21 день", "Атлас Электроника"],
    ["CAP-100U demo capacitor", "", "", "CAP-100U", "", "", "CapWorks", 10, 200, 0.63, "USD", "10 дней", "Атлас Электроника"],
])
make_book("RRFQ_Orion_Demo.xlsx", [
    ["MCU-10A demo controller", "", "", "MCU-10A", "", "", "Nordic", 1, 100, 10.90, "USD", "28 дней", "Орион Компоненты"],
    ["SENSOR-5V demo sensor", "", "", "SENSOR-5V", "", "", "Sensoria", 1, 50, "", "USD", "7 дней", "Орион Компоненты"],
    ["CAP-100U demo capacitor", "", "", "CAP-100U", "", "", "CapWorks", 10, 200, 0.71, "USD", "7 дней", "Орион Компоненты"],
])

html = '''<!doctype html><html><head><meta charset="utf-8"></head><body>
<tr data-is="ac-row" class="ac-row"><div class="ac-row-num">1</div><span class="ac-row-name">MCU-10A demo controller</span><a class="ac-item-pn-link">MCU-10A (Nordic)</a><span class="ac-part-mpq">1</span><span class="ac-item-stocks-key">120</span><span class="ac-item-offer-dlv">12 дней</span><span class="ac-item-offer-qty">100 x</span><span class="ac-item-offer-price"><price2 val="11.05" cur="USD"></price2></span><div class="ac-row-total-sum"><price2 val="1105" cur="USD"></price2></div></tr>
<tr data-is="ac-row" class="ac-row"><div class="ac-row-num">2</div><span class="ac-row-name">SENSOR-5V demo sensor</span><a class="ac-item-pn-link">SENSOR-5V (Sensoria)</a><span class="ac-part-mpq">1</span><span class="ac-item-stocks-key">0</span><span class="ac-item-offer-dlv">нет предложения</span><span class="ac-item-offer-qty">50 x</span></tr>
<tr data-is="ac-row" class="ac-row"><div class="ac-row-num">3</div><span class="ac-row-name">CAP-100U demo capacitor</span><a class="ac-item-pn-link">CAP-100U (CapWorks)</a><span class="ac-part-mpq">10</span><span class="ac-item-stocks-key">800</span><span class="ac-item-offer-dlv">9 дней</span><span class="ac-item-offer-qty">200 x</span><span class="ac-item-offer-price"><price2 val="0.60" cur="USD"></price2></span><div class="ac-row-total-sum"><price2 val="120" cur="USD"></price2></div></tr>
<div data-set="cheap"><price2 val="1225" cur="USD"></price2><span class="bom-settings-offer-type-dlv">12 дней</span></div>
<div data-set="optimal"><price2 val="1240" cur="USD"></price2><span class="bom-settings-offer-type-dlv">9 дней</span></div>
</body></html>'''
(ROOT / "SDS_Compel_DEMO.html").write_text(html, encoding="utf-8")
print(ROOT)
